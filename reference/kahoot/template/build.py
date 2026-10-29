#!/usr/bin/env python3
"""Kahoot template: kahoot-N.md -> Kahoot import .xlsx + кодын PNG зургууд.

Хэрэглээ:
    python3 kahoot/template/build.py kahoot/kahoot-3.md
    python3 kahoot/template/build.py kahoot/kahoot-3.md --time 60 --out kahoot/out

Гарц: <out>/<md нэр>/  ->  kahoot-N.xlsx  +  qNN.png (кодтой асуулт бүрд)  +  cover.png (```cover байвал)
Хэрэгтэй сан: pip install openpyxl pillow pygments
"""
import argparse
import re
import sys
from pathlib import Path

# ---- STYLE: кодын зураг (өөрчлөхгүй, бүх хичээлд ижил) -------------------
W, H = 1600, 900            # 16:9 — Kahoot-ын зургийн харьцаа
BG = "#272822"              # Monokai дэвсгэр, зураг бүхлээрээ ижил өнгө
DEFAULT_FG = "#f8f8f2"
PYGMENTS_STYLE = "monokai"
PAD = 120                   # захаас код хүртэлх зай
FONT_MAX, FONT_MIN = 56, 24  # багтахгүй бол жижигрүүлнэ
LINE_H = 1.45
FONTS = [                   # mono фонт, Кирилл дэмждэг
    "/System/Library/Fonts/Menlo.ttc",
    "/System/Library/Fonts/Monaco.ttf",
    "C:/Windows/Fonts/consola.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
    "/usr/share/fonts/dejavu/DejaVuSansMono.ttf",
]

# ---- preview: ```preview блок -> хуудасны screenshot (headless Chrome) ----
PREVIEW_CSS = """
body { margin: 0; height: 100vh; display: flex; align-items: center;
       justify-content: center; background: #ffffff; font-family: Arial, sans-serif; }
.parent { width: 1100px; height: 560px; box-sizing: border-box;
          border: 8px dashed #7c3aed; background: #f5f3ff; }
.box { width: 140px; height: 140px; color: #fff; font: bold 64px/140px Arial;
       text-align: center; border-radius: 20px; }
.box:nth-child(6n+1) { background: #ef4444; }
.box:nth-child(6n+2) { background: #f59e0b; }
.box:nth-child(6n+3) { background: #22c55e; }
.box:nth-child(6n+4) { background: #3b82f6; }
.box:nth-child(6n+5) { background: #ec4899; }
.box:nth-child(6n+6) { background: #14b8a6; }
"""
CHROMES = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "C:/Program Files/Google/Chrome/Application/chrome.exe",
    "google-chrome", "chromium", "chromium-browser",
]

# ---- flow: ```flow блок -> блок схемийн зураг (хичээлийн SVG-тэй ижил өнгө) ----
FLOW_SHAPES = {             # дүрс: (дүүргэлт, хүрээ)
    "start": ("#e8f7ee", "#2fb36d"), "end": ("#e8f7ee", "#2fb36d"),   # Эхлэл / Төгсгөл — зууван
    "io": ("#eef0fd", "#4a5bd4"),                                     # Оролт / Гаралт — налуу
    "op": ("#fff4e6", "#f29b30"),                                     # Үйлдэл — тэгш өнцөгт
    "dec": ("#fdecf3", "#d6336c"),                                    # Нөхцөл — ромбо
}
FLOW_YES, FLOW_NO = "#2fb36d", "#d6336c"   # `|` салааны Тийм (зүүн) / Үгүй (баруун)
FLOW_MARK = ("#fde68a", "#d97706")   # `op!` / `io!` — тодотгосон ("?" бөглөх) дүрс
SANS_FONTS = [
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
]
SANS_BOLD = [
    "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    "/Library/Fonts/Arial Bold.ttf",
    "C:/Windows/Fonts/arialbd.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf",
]

# ---- Kahoot хязгаар -----------------------------------------------------
LIMIT_Q, LIMIT_A = 120, 75
TIMES = (5, 10, 20, 30, 60, 90, 120, 240)
DEFAULT_TIME = 60


def clean(s):
    """Kahoot markdown-г харуулдаггүй, `<` `>` дэмждэггүй."""
    s = s.replace("`", "")
    s = re.sub(r"</?(\w+)>", lambda m: f"{m.group(1)} таг", s)
    return re.sub(r"\s+", " ", s).strip()


def parse(md):
    starts = list(re.finditer(r"<summary><b>(\d+)\.\s*(.*?)</b>", md))
    if not starts:
        sys.exit("Асуулт олдсонгүй: '<summary><b>N. ...</b>' формат хэрэгтэй")
    qs = []
    for i, m in enumerate(starts):
        end = starts[i + 1].start() if i + 1 < len(starts) else len(md)
        blk = md[m.end():end]
        code = re.search(r"```(\w+)\n(.*?)\n```", blk, re.S)
        q = re.search(r"^\*\*(.+?)\*\*\s*$", blk, re.M)
        rows = re.findall(r"^\|\s*\d+\s*\|\s*(.*?)\s*\|\s*(✅)?\s*\|\s*$", blk, re.M)
        if not q or len(rows) < 2:
            sys.exit(f"Асуулт {m.group(1)}: асуулт эсвэл хариултын хүснэгт олдсонгүй")
        correct = [k + 1 for k, r in enumerate(rows) if r[1]]
        if len(correct) != 1:
            sys.exit(f"Асуулт {m.group(1)}: зөв хариу (✅) яг 1 байх ёстой")
        qs.append({
            "n": int(m.group(1)),
            "q": clean(q.group(1)),
            "answers": [clean(r[0]) for r in rows],
            "correct": correct[0],
            "lang": code.group(1) if code else None,
            "code": code.group(2) if code else None,
        })
    return qs


def validate(qs):
    bad = []
    for x in qs:
        if len(x["q"]) > LIMIT_Q:
            bad.append(f"Q{x['n']}: асуулт {len(x['q'])} > {LIMIT_Q}")
        for a in x["answers"]:
            if len(a) > LIMIT_A:
                bad.append(f"Q{x['n']}: хариулт {len(a)} > {LIMIT_A}: {a[:30]}...")
        for t in [x["q"]] + x["answers"]:
            if "<" in t or ">" in t:
                bad.append(f"Q{x['n']}: '<' '>' тэмдэг Kahoot-д алдаа өгнө: {t[:30]}")
    if bad:
        sys.exit("\n".join(bad))


def write_xlsx(qs, path, time):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["B2"] = "Quiz template"
    hdr = ["", "Question - max 120 characters",
           "Answer 1 - max 75 characters", "Answer 2 - max 75 characters",
           "Answer 3 - max 75 characters", "Answer 4 - max 75 characters",
           "Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs",
           "Correct answer(s) - choose at least one"]
    for c, h in enumerate(hdr, 1):
        ws.cell(8, c, h)
    for r, x in enumerate(qs, 9):
        ws.cell(r, 1, r - 8)
        ws.cell(r, 2, x["q"])
        for k, a in enumerate(x["answers"][:4]):
            ws.cell(r, 3 + k, a)
        ws.cell(r, 7, time)
        ws.cell(r, 8, x["correct"])
    wb.save(path)


def find_font(custom):
    for p in ([custom] if custom else []) + FONTS:
        if p and Path(p).exists():
            return p
    sys.exit("Mono фонт олдсонгүй: --font ХАЯГ.ttf өг")


def render(code, lang, path, fontpath):
    from PIL import Image, ImageDraw, ImageFont
    from pygments.lexers import get_lexer_by_name
    from pygments.styles import get_style_by_name
    from pygments.token import Error

    style = get_style_by_name(PYGMENTS_STYLE)
    lines = [[]]
    # "???" (бөглөх зай) HTML lexer-ийг эвддэг (comment болгодог) -> түр placeholder
    for tt, val in get_lexer_by_name(lang).get_tokens(code.replace("???", "qqq")):
        val = val.replace("qqq", "???")
        col = DEFAULT_FG
        if tt not in Error:
            c = style.style_for_token(tt)["color"]
            col = "#" + c if c else DEFAULT_FG
        for j, part in enumerate(val.split("\n")):
            if j:
                lines.append([])
            if part:
                lines[-1].append((part, col))
    while lines and not lines[-1]:
        lines.pop()

    size = FONT_MAX
    while True:
        font = ImageFont.truetype(fontpath, size)
        widths = [sum(font.getlength(t) for t, _ in ln) for ln in lines]
        lh = size * LINE_H
        if (max(widths) <= W - 2 * PAD and len(lines) * lh <= H - 2 * PAD) or size <= FONT_MIN:
            break
        size -= 2

    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)
    x0 = (W - max(widths)) / 2
    y0 = (H - len(lines) * lh) / 2
    for i, ln in enumerate(lines):
        x = x0
        y = y0 + i * lh + lh / 2
        for t, col in ln:
            d.text((x, y), t, font=font, fill=col, anchor="lm")
            x += font.getlength(t)
    img.save(path)


def render_preview(html, path):
    """HTML-ийг 1600x900 хуудас болгож screenshot. .parent, .box бэлэн style-тэй."""
    import shutil
    import subprocess
    import tempfile
    chrome = next((c for c in CHROMES if Path(c).exists() or shutil.which(c)), None)
    if not chrome:
        sys.exit("Chrome олдсонгүй: preview зураг хийхэд хэрэгтэй")
    with tempfile.TemporaryDirectory() as d:
        f = Path(d) / "p.html"
        f.write_text(f"<!doctype html><meta charset='utf-8'><style>{PREVIEW_CSS}</style>{html}", encoding="utf-8")
        subprocess.run([chrome, "--headless=new", "--disable-gpu", "--hide-scrollbars",
                        f"--window-size={W},{H}", f"--screenshot={Path(path).resolve()}", f.as_uri()],
                       check=True, capture_output=True)


def first_font(paths):
    p = next((p for p in paths if Path(p).exists()), None)
    if not p:
        sys.exit(f"Фонт олдсонгүй: {paths}")
    return p


def draw_flow_shape(d, kind, box, fill, stroke, w=5):
    x0, y0, x1, y1 = box
    if kind in ("start", "end"):
        d.rounded_rectangle(box, radius=(y1 - y0) // 2, fill=fill, outline=stroke, width=w)
    elif kind == "io":
        k = 30
        d.polygon([(x0 + k, y0), (x1, y0), (x1 - k, y1), (x0, y1)], fill=fill, outline=stroke, width=w)
    elif kind == "dec":
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        d.polygon([(cx, y0), (x1, cy), (cx, y1), (x0, cy)], fill=fill, outline=stroke, width=w)
    else:
        d.rounded_rectangle(box, radius=8, fill=fill, outline=stroke, width=w)


def render_flow(src, path):
    """```flow: мөр бүр `start|io|op|dec|end[!] текст`. Дээрээс доош сумтай блок схем.
    `dec` (ромбо)-ийн дараах мөр `| io Тийм-зам | io Үгүй-зам` — хоёр салаа, дараагийн дүрс дээр нийлнэ."""
    from PIL import Image, ImageDraw, ImageFont

    def shape(txt):
        kind, _, text = txt.strip().partition(" ")
        mark = kind.endswith("!")
        kind = kind.rstrip("!")
        if kind not in FLOW_SHAPES:
            sys.exit(f"flow: '{kind}' буруу (start, io, op, dec, end)")
        return (kind, text.strip(), mark)

    rows = []   # ("one", shape) эсвэл ("two", (зүүн, баруун))
    for ln in src.strip().splitlines():
        if ln.strip().startswith("|"):
            parts = [x for x in ln.strip().split("|") if x.strip()]
            if len(parts) != 2:
                sys.exit("flow: `|` мөрөнд яг 2 салаа")
            if not rows or rows[-1][0] != "one" or rows[-1][1][0] != "dec":
                sys.exit("flow: `|` мөр `dec`-ийн дараа л байна")
            rows.append(("two", (shape(parts[0]), shape(parts[1]))))
        else:
            rows.append(("one", shape(ln)))
    n = len(rows)
    hs = lambda r: 1.5 if r[0] == "one" and r[1][0] == "dec" else 1   # ромбо өндөр
    units = sum(hs(r) for r in rows) + (n - 1) * 0.5
    sh = min(110, int((H - 2 * 40) / units))   # дүрсний өндөр
    gap = int(sh * 0.5)
    size = min(46, int(sh * 0.44))
    font = ImageFont.truetype(first_font(SANS_FONTS), size)
    bold = ImageFont.truetype(first_font(SANS_BOLD), size)
    lab = ImageFont.truetype(first_font(SANS_BOLD), int(size * 0.85))
    singles = [r[1] for r in rows if r[0] == "one" and r[1][0] != "dec"]
    sw = max(560, int(max((font.getlength(t) for _, t, _ in singles), default=0)) + 160)
    branches = [b for r in rows if r[0] == "two" for b in r[1]]
    bw = max(380, int(max((font.getlength(t) for _, t, _ in branches), default=0)) + 120)
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)

    def arrow(x, y0, y1):
        d.line((x, y0, x, y1 - 14), fill="#333", width=4)
        d.polygon([(x - 11, y1 - 16), (x + 11, y1 - 16), (x, y1)], fill="#333")

    def put(s, box):
        kind, text, mark = s
        fill, stroke = FLOW_MARK if mark else FLOW_SHAPES[kind]
        draw_flow_shape(d, kind, box, fill, stroke)
        d.text(((box[0] + box[2]) / 2, (box[1] + box[3]) / 2), text, font=bold if mark else font, fill="#222", anchor="mm")

    cx = W // 2
    y = (H - units * sh) // 2
    prev = None   # өмнөх мөрийн доод цэг(үүд): [(x, y)]
    dec_mid = None
    for kind_row, item in rows:
        h = int(hs((kind_row, item)) * sh)
        if kind_row == "one":
            if prev and len(prev) == 2:   # салаа нийлэх
                my = y - gap // 2
                for px, py in prev:
                    d.line((px, py, px, my), fill="#333", width=4)
                d.line((prev[0][0], my, prev[1][0], my), fill="#333", width=4)
                arrow(cx, my, y)
            elif prev:
                arrow(cx, prev[0][1] + 4, y - 4)
            if item[0] == "dec":
                dw = max(560, int(font.getlength(item[1]) * 1.6) + 200)
                put(item, (cx - dw // 2, y, cx + dw // 2, y + h))
                dec_mid = (cx - dw // 2, cx + dw // 2, y + h // 2)
            else:
                put(item, (cx - sw // 2, y, cx + sw // 2, y + h))
            prev = [(cx, y + h)]
        else:
            lx, rx, my = dec_mid
            bx = (min(cx - bw // 2 - 50, lx - 60), max(cx + bw // 2 + 50, rx + 60))   # салааны төв, ромбын гадна
            for (ex, bxc, label, color) in ((lx, bx[0], "Тийм", FLOW_YES), (rx, bx[1], "Үгүй", FLOW_NO)):
                d.line((ex, my, bxc, my), fill="#333", width=4)
                arrow(bxc, my, y - 4)
                d.text(((ex + bxc) / 2, my - 10), label, font=lab, fill=color, anchor="mb")
            for b, bxc in zip(item, bx):
                put(b, (bxc - bw // 2, y, bxc + bw // 2, y + h))
            prev = [(bx[0], y + h), (bx[1], y + h)]
        y += h + gap
    img.save(path)


def render_cover(src, path, monofont):
    """```cover: 1-р мөр жижиг шошго, 2-р мөр гарчиг, 3-р мөр дэд гарчиг, 4-р мөр код."""
    from PIL import Image, ImageDraw, ImageFont
    from pygments.lexers import get_lexer_by_name
    from pygments.styles import get_style_by_name
    label, title, sub, code = (src.strip().splitlines() + [""] * 4)[:4]
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)
    bold = lambda s: ImageFont.truetype(first_font(SANS_BOLD), s)
    reg = lambda s: ImageFont.truetype(first_font(SANS_FONTS), s)
    L = 130
    # шошго
    f = bold(44)
    tw = f.getlength(label)
    d.rounded_rectangle((L, 150, L + tw + 56, 230), radius=40, fill="#e6db74")
    d.text((L + 28, 190), label, font=f, fill=BG, anchor="lm")
    # гарчиг, дэд гарчиг
    d.text((L, 270), title, font=bold(150), fill=DEFAULT_FG, anchor="lt")
    d.text((L, 460), sub, font=reg(50), fill="#a6e22e", anchor="lt")
    # кодын мөр (Monokai)
    if code:
        mf = ImageFont.truetype(monofont, 52)
        style = get_style_by_name(PYGMENTS_STYLE)
        toks = [(v, "#" + (style.style_for_token(t)["color"] or "f8f8f2"))
                for t, v in get_lexer_by_name("js").get_tokens(code) if v != "\n"]
        cw = sum(mf.getlength(v) for v, _ in toks)
        d.rounded_rectangle((L, 600, L + cw + 80, 720), radius=18, fill="#1e1f1a", outline="#49483e", width=3)
        x = L + 40
        for v, c in toks:
            d.text((x, 660), v, font=mf, fill=c, anchor="lm")
            x += mf.getlength(v)
    # баруун талд жижиг блок схем
    sx0, sx1, sh, gap = 1150, 1470, 70, 40
    y = 200
    for kind in ("start", "op", "io", "end"):
        fill, stroke = FLOW_SHAPES[kind]
        draw_flow_shape(d, kind, (sx0, y, sx1, y + sh), fill, stroke, w=4)
        if kind != "end":
            ax = (sx0 + sx1) // 2
            d.line((ax, y + sh + 4, ax, y + sh + gap - 14), fill="#cfcfc2", width=4)
            d.polygon([(ax - 10, y + sh + gap - 16), (ax + 10, y + sh + gap - 16), (ax, y + sh + gap - 2)], fill="#cfcfc2")
        y += sh + gap
    img.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("md", help="kahoot-N.md")
    ap.add_argument("--time", type=int, default=DEFAULT_TIME, help=f"бүх асуултын хугацаа, сек (default {DEFAULT_TIME})")
    ap.add_argument("--out", help="гарцын хавтас (default: <md хавтас>/out)")
    ap.add_argument("--font", help="mono .ttf/.ttc хаяг")
    a = ap.parse_args()
    if a.time not in TIMES:
        sys.exit(f"--time нь {TIMES}-ийн нэг байх ёстой")

    md = Path(a.md)
    qs = parse(md.read_text(encoding="utf-8"))
    validate(qs)
    out = Path(a.out) if a.out else md.parent / "out"
    out = out / md.stem
    out.mkdir(parents=True, exist_ok=True)

    write_xlsx(qs, out / f"{md.stem}.xlsx", a.time)
    font = find_font(a.font)
    cover = re.search(r"```cover\n(.*?)\n```", md.read_text(encoding="utf-8"), re.S)
    if cover:
        render_cover(cover.group(1), out / "cover.png", font)
    for x in qs:
        if x["lang"] == "flow":
            render_flow(x["code"], out / f"q{x['n']:02d}.png")
        elif x["lang"] == "preview":
            render_preview(x["code"], out / f"q{x['n']:02d}.png")
        elif x["code"]:
            render(x["code"], x["lang"], out / f"q{x['n']:02d}.png", font)

    print(f"{out}  ({len(qs)} асуулт, {a.time} сек)")
    for x in qs:
        print(f"  Q{x['n']:>2}  зөв={x['correct']}  зураг={'схем' if x['lang'] == 'flow' else 'хуудас' if x['lang'] == 'preview' else 'код' if x['code'] else '-'}  {x['q'][:50]}")


if __name__ == "__main__":
    main()
